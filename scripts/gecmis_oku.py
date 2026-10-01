"""Önceki backlink çalışması Excel'lerinden ay ay kullanılan siteleri çıkarır ve çalışma ayına göre
son N ayın listesini üretir.

Kabul ettiği dosyalar:
  - Plugin'in ürettiği <Marka>_Backlink_Internal.xlsx / _Shared.xlsx (her ay bir sayfa, "Ekim 2026")
  - Kullanıcının elindeki eski çalışma Excel'leri (sayfa adı "Ay Yıl" olmalı; değilse --default-month ile ay verilir)

Kullanım:
    python gecmis_oku.py --files a.xlsx b.xlsx --month "Kasım 2026" [--months 6] \
        [--brand-domain dagi.com.tr] --out son6ay_excel.txt

Çıktı: --out dosyasına her satıra bir domain (son N ay, çalışma ayı hariç); ekrana ay ay döküm.
Domain, sayfadaki "Domain/Site/Mecra" başlıklı sütundan okunur; başlık yoksa tüm hücrelerdeki
domain benzeri değerler alınır. Marka domain'i ve linkli URL'ler (hedef sayfalar) hariç tutulur.
"""
import argparse
import re
from collections import defaultdict

from openpyxl import load_workbook

AYLAR = ['ocak', 'şubat', 'mart', 'nisan', 'mayıs', 'haziran', 'temmuz', 'ağustos', 'eylül', 'ekim', 'kasım', 'aralık']
FOLD = {'subat': 'şubat', 'mayis': 'mayıs', 'agustos': 'ağustos', 'eylul': 'eylül', 'kasim': 'kasım', 'aralik': 'aralık'}
DOMAIN = re.compile(r'^(?:https?://)?(?:www\.)?([a-z0-9][a-z0-9\-]*(?:\.[a-z0-9\-]+)*\.[a-z]{2,})/?$', re.I)
HEADERS = ('domain', 'site', 'site adı', 'mecra', 'url')


def month_index(text):
    """'Ekim 2026', 'ekim26', 'Ekim24', 'Mart 2026 ' -> yıl*12+ay; bulunamazsa None."""
    t = text.strip().lower().replace('ı', 'ı')
    m = re.match(r'([a-zçğıöşü]+)\s*(\d{2,4})', t)
    if not m:
        return None
    name, year = m.group(1), m.group(2)
    name = FOLD.get(name, name)
    if name not in AYLAR:
        return None
    year = int(year) + (2000 if len(year) == 2 else 0)
    return year * 12 + AYLAR.index(name)


def label(idx):
    return f'{AYLAR[idx % 12].capitalize()} {idx // 12}'


def domains_in_sheet(ws, brand):
    found = set()
    cols = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        heads = [i for i, c in enumerate(row) if isinstance(c.value, str) and c.value.strip().lower() in HEADERS]
        if heads:
            cols = heads
            break
    for row in ws.iter_rows():
        cells = [row[i] for i in cols if i < len(row)] if cols else row
        for c in cells:
            v = c.value
            if not isinstance(v, str) or v.startswith('='):
                continue
            m = DOMAIN.match(v.strip())
            if m:
                d = m.group(1).lower()
                if brand and d.endswith(brand):
                    continue
                found.add(d)
    return found


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--files', nargs='+', required=True)
    ap.add_argument('--month', required=True, help='Çalışma ayı, ör. "Kasım 2026"')
    ap.add_argument('--months', type=int, default=6)
    ap.add_argument('--brand-domain', default='')
    ap.add_argument('--default-month', help='Sayfa adı ay içermiyorsa kullanılacak ay')
    ap.add_argument('--out', default='son6ay_excel.txt')
    a = ap.parse_args()
    cur = month_index(a.month)
    brand = a.brand_domain.lower().replace('www.', '')
    by_month, unknown = defaultdict(set), defaultdict(set)
    for f in a.files:
        wb = load_workbook(f, read_only=False)
        for ws in wb.worksheets:
            idx = month_index(ws.title) or (month_index(a.default_month) if a.default_month else None)
            ds = domains_in_sheet(ws, brand)
            if idx is None:
                unknown[f'{f} :: {ws.title}'] |= ds
            else:
                by_month[idx] |= ds
    window = [i for i in by_month if cur - a.months <= i < cur]
    keep = sorted({d for i in window for d in by_month[i]})
    open(a.out, 'w', encoding='utf-8').write('\n'.join(keep) + '\n')
    for i in sorted(by_month):
        tag = 'son ' + str(a.months) + ' ay' if i in window else ('çalışma ayı' if i == cur else 'kapsam dışı')
        print(f'{label(i):14} ({tag}): {", ".join(sorted(by_month[i]))}')
    for k, ds in unknown.items():
        print(f'⚠ Ayı anlaşılamayan sayfa "{k}": {", ".join(sorted(ds)) or "domain yok"} (kullanıcıya hangi ay olduğunu sor)')
    print(f'→ {len(keep)} domain son {a.months} ay listesine yazıldı: {a.out}')
